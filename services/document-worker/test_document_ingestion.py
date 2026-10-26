import io
import unittest
import zipfile
from openpyxl import Workbook
from document_ingestion import parse_document

class BulkImportTests(unittest.TestCase):
    def test_csv(self):
        items=parse_document('plots.csv',b'title,property_type,price\nGreen Court,land,4000000\n')
        self.assertEqual(items[0].rows[0]['title'],'Green Court')
        self.assertTrue(items[0].review_required)

    def test_spreadsheet(self):
        workbook=Workbook()
        sheet=workbook.active
        sheet.append(['title','property_type','price'])
        sheet.append(['Landshop Estate','land','1250000'])
        buffer=io.BytesIO();workbook.save(buffer)
        results=parse_document('plots.xlsx',buffer.getvalue())
        self.assertEqual(results[0].rows[0]['price'],'1250000')

    def test_xml(self):
        data=b'<properties><property><title>Green Court</title><property_type>land</property_type></property></properties>'
        self.assertEqual(parse_document('data.xml',data)[0].rows[0]['title'],'Green Court')

    def test_zip_path_traversal(self):
        buf=io.BytesIO()
        with zipfile.ZipFile(buf,'w') as z:
            z.writestr('../escape.csv','title,property_type\nABC,land')
        with self.assertRaises(ValueError):
            parse_document('imports.zip',buf.getvalue())

    def test_zip_nested(self):
        buf=io.BytesIO()
        with zipfile.ZipFile(buf,'w') as z:
            z.writestr('nested.zip','not a proper archive')
        with self.assertRaises(ValueError):
            parse_document('imports.zip',buf.getvalue())

    def test_docx_does_not_auto_publish(self):
        from docx import Document
        doc=Document();doc.add_paragraph('A verified estate listing is required.')
        buffer=io.BytesIO();doc.save(buffer)
        result=parse_document('brief.docx',buffer.getvalue())[0]
        self.assertEqual(result.rows,[])
        self.assertIn('estate',result.review_text)

if __name__=='__main__':
    unittest.main()
