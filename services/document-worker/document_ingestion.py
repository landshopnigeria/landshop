"""Untrusted bulk-import parser. Parsing ONLY: never writes or publishes properties.

Supports CSV, XLSX, XML, DOCX, PDF and ZIP of supported inputs.
PDF/DOCX are returned as review-required extracted text, not invented listings.
ZIP traversal, encrypted entries, nested ZIPs and archive bombs are rejected.
"""
from __future__ import annotations
import csv
import io
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import PurePosixPath
from typing import Any
from defusedxml import ElementTree as SafeET
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

MAX_INPUT_BYTES = 20 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 40 * 1024 * 1024
MAX_ROWS = 25000
MAX_ARCHIVE_ENTRIES = 100
MAX_PDF_PAGES = 150
SUPPORTED = {'.csv', '.xlsx', '.xml', '.pdf', '.docx', '.zip'}
FIELDS = {'title','estate','state','city','property_type','price','plot_number','size_sqm'}

@dataclass
class Extraction:
    filename: str
    rows: list[dict[str,str]] = field(default_factory=list)
    review_text: str = ''
    warnings: list[str] = field(default_factory=list)
    review_required: bool = True

def _suffix(name: str) -> str:
    return '.' + name.rsplit('.', 1)[-1].lower() if '.' in name else ''

def _clean(value: Any) -> str:
    if value is None: return ''
    return str(value).strip()[:2000]

def _rows_from_dicts(items, name: str) -> Extraction:
    rows: list[dict[str,str]] = []
    warnings: list[str] = []
    for idx, raw in enumerate(items, 1):
        if idx > MAX_ROWS: raise ValueError('Too many import records')
        if not isinstance(raw, dict): raise ValueError('Import row must be an object')
        values = {str(k).strip().lower():_clean(v) for k,v in raw.items()}
        unknown = set(values) - FIELDS
        if unknown: warnings.append('Row %d: unknown columns %s' % (idx, ', '.join(sorted(unknown))))
        rows.append({k:values.get(k,'') for k in FIELDS if k in values})
    return Extraction(filename=name, rows=rows, warnings=warnings[:100])

def _csv(payload: bytes, name: str) -> Extraction:
    source=io.StringIO(payload.decode('utf-8-sig'), newline='')
    reader=csv.DictReader(source)
    if not reader.fieldnames: raise ValueError('Missing CSV column headings')
    if len(reader.fieldnames)!=len(set(reader.fieldnames)):
        raise ValueError('Duplicate CSV headings')
    return _rows_from_dicts(reader,name)

def _validate_office_archive(payload: bytes) -> None:
    """Prevent large or path-traversing internal Office ZIP members."""
    try:
        with zipfile.ZipFile(io.BytesIO(payload)) as office:
            entries=office.infolist()
            if len(entries)>500: raise ValueError('Office document contains too many parts')
            total=0
            for entry in entries:
                path=PurePosixPath(entry.filename.replace(chr(92),'/'))
                if path.is_absolute() or '..' in path.parts:
                    raise ValueError('Unsafe Office document member')
                if entry.flag_bits&1: raise ValueError('Encrypted Office document not allowed')
                total+=entry.file_size
                if total>MAX_UNCOMPRESSED_BYTES:
                    raise ValueError('Expanded Office document is too large')
                if entry.file_size and entry.compress_size and entry.file_size/entry.compress_size>100:
                    raise ValueError('Office archive expansion ratio exceeded')
    except zipfile.BadZipFile as exc:
        raise ValueError('Malformed Office document') from exc

def _xlsx(payload: bytes, name: str) -> Extraction:
    _validate_office_archive(payload)
    workbook=load_workbook(io.BytesIO(payload),read_only=True,data_only=True)
    try:
        sheet=workbook.active
        if sheet is None: raise ValueError('Workbook has no active sheet')
        iterator=sheet.iter_rows(values_only=True)
        header=next(iterator,None)
        if not header: raise ValueError('Worksheet is empty')
        keys=[_clean(k).lower() for k in header]
        if len(keys)!=len(set(keys)): raise ValueError('Duplicate spreadsheet headings')
        def entries():
            for row in iterator:
                if not any(value is not None for value in row):continue
                yield dict(zip(keys,row))
        return _rows_from_dicts(entries(),name)
    finally: workbook.close()

def _xml(payload: bytes, name: str) -> Extraction:
    root=SafeET.fromstring(payload)
    if root.tag.lower() not in ('properties','listings'):
        raise ValueError('Expected a properties or listings XML root')
    def entries():
        for child in root:
            if child.tag.lower() not in ('property','listing'):
                raise ValueError('Unexpected XML element')
            yield {element.tag:element.text for element in child}
    return _rows_from_dicts(entries(),name)

def _docx(payload: bytes,name: str) -> Extraction:
    _validate_office_archive(payload)
    document=Document(io.BytesIO(payload))
    text='\n'.join(p.text for p in document.paragraphs if p.text.strip())
    for table in document.tables:
        text+='\n'+'\n'.join(' | '.join(c.text for c in row.cells) for row in table.rows)
    return Extraction(name,review_text=text[:40000],warnings=['DOCX extracted for human field mapping; no automatic listing created'])

def _pdf(payload: bytes,name: str) -> Extraction:
    reader=PdfReader(io.BytesIO(payload),strict=True)
    if reader.is_encrypted: raise ValueError('Encrypted PDF not accepted')
    if len(reader.pages)>MAX_PDF_PAGES: raise ValueError('PDF page limit exceeded')
    text='\n'.join((page.extract_text() or '') for page in reader.pages)
    return Extraction(name,review_text=text[:40000],warnings=['PDF text extraction requires human mapping; scanned PDFs may contain no text'])

def parse_document(filename: str,payload: bytes) -> list[Extraction]:
    if not isinstance(payload,bytes) or not payload or len(payload)>MAX_INPUT_BYTES:
        raise ValueError('File is missing or exceeds 20 MiB')
    ext=_suffix(filename)
    if ext not in SUPPORTED: raise ValueError('Unsupported file format')
    if ext=='.csv':return [_csv(payload,filename)]
    if ext=='.xlsx':return [_xlsx(payload,filename)]
    if ext=='.xml':return [_xml(payload,filename)]
    if ext=='.docx':return [_docx(payload,filename)]
    if ext=='.pdf':return [_pdf(payload,filename)]
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        entries=archive.infolist()
        if len(entries)>MAX_ARCHIVE_ENTRIES: raise ValueError('Too many ZIP entries')
        total=0
        result=[]
        for entry in entries:
            if entry.is_dir():continue
            path=PurePosixPath(entry.filename.replace('\\','/'))
            if path.is_absolute() or '..' in path.parts or len(path.parts)>5:
                raise ValueError('Unsafe archive path')
            if entry.flag_bits & 1: raise ValueError('Encrypted archive entries are not supported')
            if entry.file_size>MAX_INPUT_BYTES: raise ValueError('Oversized archive entry')
            total+=entry.file_size
            if total>MAX_UNCOMPRESSED_BYTES: raise ValueError('Archive extraction limit exceeded')
            if entry.file_size and entry.compress_size and entry.file_size/entry.compress_size>100:
                raise ValueError('Suspicious archive expansion ratio')
            suffix=_suffix(entry.filename)
            if suffix in ('.zip','.exe','.js','.html','.svg'):
                raise ValueError('Unsupported executable or nested archive format')
            if suffix not in SUPPORTED:
                continue
            data=archive.read(entry)
            result.extend(parse_document(entry.filename,data))
        if not result: raise ValueError('No supported files found in ZIP')
        return result
